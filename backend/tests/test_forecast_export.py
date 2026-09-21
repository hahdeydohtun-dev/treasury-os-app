from httpx import AsyncClient
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.rbac import User


async def _login(client: AsyncClient, user: User) -> str:
    resp = await client.post(
        "/api/v1/auth/login", json={"email": user.email, "password": "Password123!"}
    )
    assert resp.status_code == 200
    return resp.json()["access_token"]


async def test_export_returns_valid_xlsx_with_expected_sheets(
    client: AsyncClient, db_session: AsyncSession, group_wide_manager: User,
    demo_group_and_entities, forecast_categories,
):
    import io

    import openpyxl

    group, entity_a, _ = demo_group_and_entities
    await db_session.commit()
    token = await _login(client, group_wide_manager)
    headers = {"Authorization": f"Bearer {token}"}

    create_resp = await client.post(
        "/api/v1/forecast",
        json={
            "legal_entity_id": str(entity_a.id), "forecast_start_date": "2026-06-01",
            "reporting_currency_code": "NGN",
        },
        headers=headers,
    )
    forecast_id = create_resp.json()["id"]
    await client.post(f"/api/v1/forecast/{forecast_id}/calculate", headers=headers)

    export_resp = await client.get(f"/api/v1/forecast/{forecast_id}/export", headers=headers)
    assert export_resp.status_code == 200
    assert export_resp.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    wb = openpyxl.load_workbook(io.BytesIO(export_resp.content))
    assert "Weekly Summary" in wb.sheetnames
    assert "Cash Flow Lines" in wb.sheetnames
    ws = wb["Weekly Summary"]
    assert ws.max_row == 14


async def test_export_requires_calculated_forecast(
    client: AsyncClient, db_session: AsyncSession, group_wide_manager: User,
    demo_group_and_entities, forecast_categories,
):
    group, entity_a, _ = demo_group_and_entities
    await db_session.commit()
    token = await _login(client, group_wide_manager)
    headers = {"Authorization": f"Bearer {token}"}

    create_resp = await client.post(
        "/api/v1/forecast",
        json={
            "legal_entity_id": str(entity_a.id), "forecast_start_date": "2026-06-01",
            "reporting_currency_code": "NGN",
        },
        headers=headers,
    )
    forecast_id = create_resp.json()["id"]

    export_resp = await client.get(f"/api/v1/forecast/{forecast_id}/export", headers=headers)
    assert export_resp.status_code == 400
