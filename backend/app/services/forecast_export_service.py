"""
13-Week Forecast Excel export (SECTION 48).

Builds a multi-sheet workbook directly from the same data every API
endpoint reads - no separate "reporting" copy of the numbers, so the
export can never drift from what the dashboard shows. Uses openpyxl
(already a dependency for the Excel Data Hub).
"""
import io

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.entity import LegalEntity
from app.models.forecast import Forecast, ForecastLine
from app.services.forecast_variance_service import compute_accuracy, compute_variance

BOLD = Font(bold=True)


def _autosize(ws) -> None:
    for col_idx, column_cells in enumerate(ws.columns, start=1):
        length = max((len(str(c.value)) if c.value is not None else 0) for c in column_cells)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(length + 2, 10), 40)


async def build_forecast_export(db: AsyncSession, forecast: Forecast) -> bytes:
    weeks = sorted(forecast.weeks, key=lambda w: w.week_number)
    wb = openpyxl.Workbook()

    meta = wb.active
    meta.title = "Forecast Metadata"
    rows = [
        ("Forecast ID", str(forecast.id)),
        ("Group ID", str(forecast.group_id) if forecast.group_id else ""),
        ("Legal Entity ID", str(forecast.legal_entity_id) if forecast.legal_entity_id else "GROUP-WIDE"),
        ("Forecast Start Date", str(forecast.forecast_start_date)),
        ("Forecast End Date", str(forecast.forecast_end_date)),
        ("Scenario", forecast.scenario.value),
        ("Value Basis", forecast.value_basis.value),
        ("Reporting Currency", forecast.reporting_currency_code),
        ("Version", forecast.version),
        ("Status", forecast.status.value),
        ("Assumptions Notes", forecast.assumptions_notes or ""),
    ]
    for label, value in rows:
        meta.append([label, value])
    for warning in forecast.data_quality_warnings:
        meta.append(["Data Quality Warning", warning])
    meta["A1"].font = BOLD
    _autosize(meta)

    summary_ws = wb.create_sheet("Weekly Summary")
    headers = [
        "Week", "Start Date", "End Date", "Status", "Opening Cash", "Inflows", "Outflows",
        "Net Transfers", "Net Cash Flow", "Closing Cash", "Minimum Liquidity",
        "Surplus / Gap", "Liquidity Coverage",
    ]
    summary_ws.append(headers)
    for cell in summary_ws[1]:
        cell.font = BOLD
    for w in weeks:
        summary_ws.append([
            w.week_number, w.start_date, w.end_date, w.status.value, float(w.opening_cash),
            float(w.total_inflows), float(w.total_outflows), float(w.net_transfers),
            float(w.net_cash_flow), float(w.closing_cash), float(w.minimum_required_liquidity),
            float(w.surplus_or_gap),
            float(w.liquidity_coverage_ratio) if w.liquidity_coverage_ratio is not None else None,
        ])
    _autosize(summary_ws)

    lines_ws = wb.create_sheet("Cash Flow Lines")
    line_headers = [
        "Week", "Entity ID", "Category", "Direction", "Currency", "Original Amount",
        "Probability %", "Adjusted Amount", "Reporting Currency", "Reporting Amount",
        "FX Rate", "Source Type", "Source ID", "Description", "Counterparty",
    ]
    lines_ws.append(line_headers)
    for cell in lines_ws[1]:
        cell.font = BOLD

    week_number_by_id = {w.id: w.week_number for w in weeks}
    lines_result = await db.execute(select(ForecastLine).where(ForecastLine.forecast_id == forecast.id))
    all_lines = list(lines_result.scalars().all())
    for line in all_lines:
        lines_ws.append([
            week_number_by_id.get(line.week_id), str(line.legal_entity_id), line.category_code,
            line.direction.value, line.transaction_currency_code, float(line.original_amount),
            line.probability, float(line.adjusted_amount), line.reporting_currency_code,
            float(line.reporting_amount), float(line.fx_rate) if line.fx_rate else None,
            line.source_type.value, line.source_id, line.description, line.counterparty,
        ])
    _autosize(lines_ws)

    entity_ws = wb.create_sheet("Entity Breakdown")
    entity_ws.append(["Entity", "Total Inflows", "Total Outflows", "Net Cash Flow"])
    for cell in entity_ws[1]:
        cell.font = BOLD
    entity_ids = {line.legal_entity_id for line in all_lines}
    entities_result = await db.execute(select(LegalEntity).where(LegalEntity.id.in_(entity_ids)))
    entity_names = {e.id: e.name for e in entities_result.scalars().all()}
    for entity_id in entity_ids:
        entity_lines = [ln for ln in all_lines if ln.legal_entity_id == entity_id]
        inflows = sum((ln.reporting_amount for ln in entity_lines if ln.direction.value == "INFLOW"), 0)
        outflows = sum((ln.reporting_amount for ln in entity_lines if ln.direction.value == "OUTFLOW"), 0)
        entity_ws.append([
            entity_names.get(entity_id, str(entity_id)), float(inflows), float(outflows),
            float(inflows - outflows),
        ])
    _autosize(entity_ws)

    currency_ws = wb.create_sheet("Currency Breakdown")
    currency_ws.append(["Currency", "Total Inflows (native)", "Total Outflows (native)"])
    for cell in currency_ws[1]:
        cell.font = BOLD
    currencies = {line.transaction_currency_code for line in all_lines}
    for currency in currencies:
        currency_lines = [ln for ln in all_lines if ln.transaction_currency_code == currency]
        inflows = sum((ln.adjusted_amount for ln in currency_lines if ln.direction.value == "INFLOW"), 0)
        outflows = sum((ln.adjusted_amount for ln in currency_lines if ln.direction.value == "OUTFLOW"), 0)
        currency_ws.append([currency, float(inflows), float(outflows)])
    _autosize(currency_ws)

    gaps_ws = wb.create_sheet("Funding Gaps & Surplus")
    gaps_ws.append(["Week", "Closing Cash", "Minimum Liquidity", "Surplus / Gap"])
    for cell in gaps_ws[1]:
        cell.font = BOLD
    for w in weeks:
        if w.surplus_or_gap != 0:
            gaps_ws.append([
                w.week_number, float(w.closing_cash), float(w.minimum_required_liquidity),
                float(w.surplus_or_gap),
            ])
    _autosize(gaps_ws)

    variance_ws = wb.create_sheet("Forecast vs Actual")
    variance_ws.append(["Week", "Forecast Amount", "Actual Amount", "Variance", "Variance %"])
    for cell in variance_ws[1]:
        cell.font = BOLD
    variance_rows = await compute_variance(db, forecast, group_by="week")
    for row in variance_rows:
        variance_ws.append([
            row.week_number, float(row.forecast_amount), float(row.actual_amount),
            float(row.variance_amount),
            float(row.variance_percentage) if row.variance_percentage is not None else None,
        ])
    _autosize(variance_ws)

    accuracy_ws = wb.create_sheet("Accuracy")
    accuracy = await compute_accuracy(db, forecast, target_variance_pct=5)
    accuracy_ws.append(["Metric", "Value"])
    for cell in accuracy_ws[1]:
        cell.font = BOLD
    for key, value in accuracy.items():
        accuracy_ws.append([key, str(value)])
    _autosize(accuracy_ws)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
